#!/usr/bin/env python3
"""Convert DOCX documents to structured Markdown through Pandoc.

DOCX is preprocessed structurally: Word Heading 1-6 styles remain unchanged,
monospaced paragraphs become Source Code blocks, and layout-only page indents
are removed.  Pandoc then preserves lists, links, tables, anchored images and
document order.  A Lua filter merges adjacent code paragraphs into one fence.
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import uuid
import xml.etree.ElementTree as ET
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator
from urllib.parse import unquote
from zipfile import BadZipFile, ZipFile


WORD_SUFFIXES = {".docx"}
MONOSPACED_FONTS = {
    "courier",
    "courier new",
    "consolas",
    "lucida console",
    "menlo",
    "monaco",
}
HEADING_STYLE_RE = re.compile(r"^Heading\s+([1-6])$", re.IGNORECASE)
SENTENCE_END_RE = re.compile(r"[.!?。！？:;…](?:[\"'”’）)\]]*)$")
COMMAND_OPTION_RE = re.compile(
    r"(?<![\w-])(?:--?|\u2010|\u2011|\u2212|\ufe63|\uff0d)[A-Za-z_][\w-]*"
)
QUOTED_COMMAND_RE = re.compile(
    r"[\"“][^\"”\r\n]*(?<![\w-])(?:--?|\u2010|\u2011|\u2212|\ufe63|\uff0d)[A-Za-z_][\w-]*[^\"”\r\n]*[\"”]"
)
class WordConversionError(RuntimeError):
    """Raised when a DOCX document cannot be converted safely."""


@dataclass
class WordConversionReport:
    source: str
    output: str
    pages: int = 0
    text_characters: int = 0
    headings: dict[str, int] = field(default_factory=lambda: defaultdict(int))
    images: int = 0
    tables: int = 0
    code_blocks: int = 0
    bookmarks_total: int = 0
    bookmarks_matched: int = 0
    unmatched_bookmarks: list[dict[str, object]] = field(default_factory=list)
    ocr_used: bool = False
    conversion_engine: str = "pandoc-docx"
    native_text_rejection: str | None = None
    chunks: int = 1

    def as_dict(self) -> dict[str, object]:
        return {
            "source": self.source,
            "output": self.output,
            "pages": self.pages,
            "text_characters": self.text_characters,
            "headings": dict(self.headings),
            "images": self.images,
            "tables": self.tables,
            "code_blocks": self.code_blocks,
            "bookmarks_total": self.bookmarks_total,
            "bookmarks_matched": self.bookmarks_matched,
            "unmatched_bookmarks": self.unmatched_bookmarks,
            "ocr_used": self.ocr_used,
            "conversion_engine": self.conversion_engine,
            "native_text_rejection": self.native_text_rejection,
            "chunks": self.chunks,
        }


def safe_name(value: str, fallback: str = "document") -> str:
    value = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", value).strip(" .")
    return value or fallback


def find_pandoc() -> Path:
    executable = shutil.which("pandoc")
    if executable:
        return Path(executable).resolve()
    try:
        import pypandoc  # type: ignore[import-not-found]

        bundled = Path(pypandoc.get_pandoc_path())
        if bundled.is_file():
            return bundled.resolve()
    except (ImportError, OSError, RuntimeError):
        pass
    raise WordConversionError(
        "未找到 Pandoc。请运行 pip install -r requirements.txt，"
        "或从 https://pandoc.org/installing.html 安装 Pandoc。"
    )


def run_checked(command: list[str], *, cwd: Path | None = None) -> None:
    completed = subprocess.run(
        command,
        cwd=str(cwd) if cwd else None,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if completed.returncode:
        details = (completed.stderr or completed.stdout).strip()
        raise WordConversionError(details or f"外部转换命令失败，退出码 {completed.returncode}")


def iter_document_paragraphs(document: object) -> Iterator[object]:
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    for element in document.element.body.iter(qn("w:p")):  # type: ignore[attr-defined]
        yield Paragraph(element, document)


def paragraph_heading_level(paragraph: object) -> int | None:
    match = HEADING_STYLE_RE.match(str(paragraph.style.name))  # type: ignore[attr-defined]
    return int(match.group(1)) if match else None


def monospaced_ratio(paragraph: object) -> float:
    total = 0
    monospaced = 0
    style_font = str(paragraph.style.font.name or "").casefold()  # type: ignore[attr-defined]
    for run in paragraph.runs:  # type: ignore[attr-defined]
        count = len(re.sub(r"\s+", "", run.text))
        if not count:
            continue
        total += count
        font = str(run.font.name or style_font).casefold()
        if font in MONOSPACED_FONTS:
            monospaced += count
    return monospaced / total if total else 0.0


def merge_split_body_paragraphs(document: object) -> int:
    """Join body paragraphs that Word split only for page layout.

    Generated manuals sometimes store one sentence as two ``Body Text``
    paragraphs, optionally separated by an empty page-break paragraph.  A
    fragment without sentence-ending punctuation is safely continued by the
    following Body Text paragraph.  Code paragraphs are excluded because
    their line boundaries are meaningful.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraphs = list(iter_document_paragraphs(document))
    merged = 0
    index = 0
    while index < len(paragraphs):
        previous = paragraphs[index]
        previous_text = str(previous.text).strip()
        if (
            not previous_text
            or str(previous.style.name) != "Body Text"
            or SENTENCE_END_RE.search(previous_text)
            or monospaced_ratio(previous) >= 0.50
        ):
            index += 1
            continue

        parent = previous._p.getparent()
        cursor = index + 1
        while cursor < len(paragraphs):
            empty_paragraphs: list[object] = []
            while cursor < len(paragraphs):
                candidate = paragraphs[cursor]
                if candidate._p.getparent() is not parent or str(candidate.text).strip():
                    break
                empty_paragraphs.append(candidate)
                cursor += 1
            if cursor >= len(paragraphs):
                break
            following = paragraphs[cursor]
            if (
                following._p.getparent() is not parent
                or str(following.style.name) != "Body Text"
                or not str(following.text).strip()
                or monospaced_ratio(following) >= 0.50
            ):
                break

            spacer_run = OxmlElement("w:r")
            spacer_text = OxmlElement("w:t")
            spacer_text.set(qn("xml:space"), "preserve")
            spacer_text.text = " "
            spacer_run.append(spacer_text)
            previous._p.append(spacer_run)
            for child in list(following._p):
                if child.tag != qn("w:pPr"):
                    previous._p.append(child)
            for empty in empty_paragraphs:
                parent.remove(empty._p)
            parent.remove(following._p)
            merged += 1
            cursor += 1
            if SENTENCE_END_RE.search(str(previous.text).strip()):
                break

        # ``cursor`` points to the next paragraph not merged into ``previous``.
        index = max(index + 1, cursor)
    return merged


def normalize_command_placeholder_runs(document: object) -> int:
    """Keep command syntax out of Markdown emphasis and math parsing.

    Generated manuals often render command metavariables with an italic
    character style.  Pandoc faithfully writes those runs as ``*length*`` and
    similar fragments.  Around bracketed command options, some Markdown
    renderers then treat the fragments as mathematics and can display one
    character per line.

    Detection is based on generic command-line option syntax rather than a
    product-specific command name.  A paragraph is considered command-bearing
    when it contains multiple options, or a quoted command with at least one
    option.  Only italic runs in such paragraphs lose their visual-only
    emphasis.  Typographic option hyphens are also changed to ASCII.
    """
    changed = 0
    for paragraph in iter_document_paragraphs(document):
        paragraph_text = str(paragraph.text)
        options = COMMAND_OPTION_RE.findall(paragraph_text)
        if len(options) < 2 and not QUOTED_COMMAND_RE.search(paragraph_text):
            continue
        for run in paragraph.runs:
            normalized = re.sub(
                r"[\u2010\u2011\u2212\ufe63\uff0d](?=[A-Za-z_])",
                "-",
                str(run.text),
            )
            if normalized != run.text:
                run.text = normalized
                changed += 1
            if run.italic:
                run.italic = False
                changed += 1
    return changed


def paragraph_indent_inches(paragraph: object) -> float | None:
    indent = paragraph.paragraph_format.left_indent  # type: ignore[attr-defined]
    return float(indent.inches) if indent is not None else None


def prepend_spaces(paragraph: object, count: int) -> None:
    if count <= 0 or str(paragraph.text).startswith((" ", "\t")):  # type: ignore[attr-defined]
        return
    for run in paragraph.runs:  # type: ignore[attr-defined]
        if run.text:
            run.text = " " * count + run.text
            return
    paragraph.add_run(" " * count)  # type: ignore[attr-defined]


def preprocess_docx(source: Path, output_docx: Path) -> tuple[dict[str, int], int, int, int]:
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Pt
    except ImportError as exc:
        raise WordConversionError(
            "缺少 python-docx；请运行 pip install -r requirements.txt。"
        ) from exc

    try:
        document = Document(str(source))
    except (ValueError, KeyError, BadZipFile) as exc:
        raise WordConversionError(f"无法读取 Word OOXML 文档：{source}：{exc}") from exc

    headings: dict[str, int] = defaultdict(int)
    for paragraph in document.paragraphs:
        level = paragraph_heading_level(paragraph)
        if level is not None and paragraph.text.strip():
            headings[str(level)] += 1

    merge_split_body_paragraphs(document)
    normalize_command_placeholder_runs(document)

    style_names = {style.name for style in document.styles}
    code_style = (
        document.styles["Source Code"]
        if "Source Code" in style_names
        else document.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    )
    code_style.font.name = "Courier New"
    code_style.font.size = Pt(9)

    paragraphs = list(iter_document_paragraphs(document))
    code_flags = [
        bool(str(paragraph.text).strip())
        and paragraph_heading_level(paragraph) is None
        and (
            str(paragraph.style.name).casefold() in {"source code", "code", "preformatted text"}
            or monospaced_ratio(paragraph) >= 0.70
        )
        for paragraph in paragraphs
    ]

    index = 0
    while index < len(paragraphs):
        if not code_flags[index]:
            index += 1
            continue
        end = index + 1
        while end < len(paragraphs) and code_flags[end]:
            end += 1
        indents = [
            value
            for value in (paragraph_indent_inches(item) for item in paragraphs[index:end])
            if value is not None
        ]
        baseline = min(indents) if indents else 0.0
        for paragraph in paragraphs[index:end]:
            indent = paragraph_indent_inches(paragraph)
            if indent is not None:
                relative = max(0.0, indent - baseline)
                prepend_spaces(paragraph, min(24, round(relative / 0.12) * 2))
            paragraph.style = code_style
            paragraph.paragraph_format.left_indent = None
            paragraph.paragraph_format.right_indent = None
            paragraph.paragraph_format.first_line_indent = None
        index = end

    for paragraph, is_code in zip(paragraphs, code_flags):
        if is_code:
            continue
        style_name = str(paragraph.style.name)
        indent = paragraph_indent_inches(paragraph)
        has_numbering = (
            paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
        )
        # Many generated manuals use a 0.5-0.7 inch page-layout offset on all
        # body paragraphs. Pandoc interprets that as a quotation. Clear only
        # this layout offset; real list numbering and heading styles remain.
        if (
            style_name in {"Body Text", "Normal"}
            and not has_numbering
            and indent is not None
            and indent >= 0.30
        ):
            paragraph.paragraph_format.left_indent = None
            paragraph.paragraph_format.right_indent = None

    document.save(str(output_docx))
    return dict(headings), len(document.tables), len(paragraphs), sum(code_flags)


def word_page_count(docx_path: Path) -> int:
    try:
        with ZipFile(docx_path) as archive:
            root = ET.fromstring(archive.read("docProps/app.xml"))
        for node in root.iter():
            if node.tag.rsplit("}", 1)[-1] == "Pages" and node.text:
                return int(node.text)
    except (BadZipFile, KeyError, ET.ParseError, TypeError, ValueError):
        pass
    return 0


def flatten_images(markdown: str, output_dir: Path) -> tuple[str, int]:
    image_dir = output_dir / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    files = sorted((path for path in image_dir.rglob("*") if path.is_file()), key=str)
    used_names: set[str] = set()
    replacements: list[tuple[str, str]] = []
    for source in files:
        old_relative = source.relative_to(output_dir).as_posix()
        candidate = safe_name(source.name, "image.bin")
        stem = Path(candidate).stem
        suffix = Path(candidate).suffix
        serial = 2
        while candidate.casefold() in used_names:
            candidate = f"{stem}-{serial}{suffix}"
            serial += 1
        used_names.add(candidate.casefold())
        destination = image_dir / candidate
        if source != destination:
            source.replace(destination)
        new_relative = f"images/{candidate}"
        replacements.append((old_relative, new_relative))

    for directory in sorted(
        (path for path in image_dir.rglob("*") if path.is_dir()),
        key=lambda path: len(path.parts),
        reverse=True,
    ):
        try:
            directory.rmdir()
        except OSError:
            pass

    for old, new in replacements:
        backslash_old = old.replace("/", "\\")
        variants = {
            old,
            backslash_old,
            f"./{old}",
            ".\\" + backslash_old,
        }
        for variant in variants:
            markdown = markdown.replace(variant, new)
    return markdown, len(files)


def count_markdown_headings(markdown: str) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    fence: str | None = None
    for line in markdown.splitlines():
        stripped = line.lstrip()
        marker_match = re.match(r"^(`{3,}|~{3,})", stripped)
        if marker_match:
            marker = marker_match.group(1)[0]
            fence = None if fence == marker else marker
            continue
        if fence is not None:
            continue
        match = re.match(r"^(#{1,6})\s+\S", line)
        if match:
            counts[str(len(match.group(1)))] += 1
    return dict(counts)


def count_code_blocks(markdown: str) -> int:
    opened = 0
    fence: str | None = None
    for line in markdown.splitlines():
        match = re.match(r"^(`{3,}|~{3,})", line.lstrip())
        if not match:
            continue
        marker = match.group(1)[0]
        if fence is None:
            fence = marker
            opened += 1
        elif fence == marker:
            fence = None
    return opened


def validate_local_images(markdown: str, output_dir: Path) -> None:
    references = re.findall(r"!\[[^\]]*\]\((?:<)?([^)>\s]+)", markdown)
    references.extend(re.findall(r"<img\b[^>]*\bsrc=[\"']([^\"']+)", markdown, re.IGNORECASE))
    missing: list[str] = []
    for reference in references:
        decoded = unquote(reference).replace("\\", "/")
        if re.match(r"^(?:https?:|data:)", decoded, re.IGNORECASE):
            continue
        candidate = (output_dir / decoded).resolve()
        try:
            candidate.relative_to(output_dir.resolve())
        except ValueError:
            missing.append(reference)
            continue
        if not candidate.is_file():
            missing.append(reference)
    if missing:
        preview = ", ".join(missing[:5])
        raise WordConversionError(f"Markdown 中存在缺失的图片引用：{preview}")


def atomic_replace_dir(source: Path, destination: Path, overwrite: bool) -> None:
    if destination.exists():
        if not overwrite:
            raise WordConversionError(f"输出目录已存在：{destination}；使用 --overwrite 可覆盖。")
        shutil.rmtree(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    source.replace(destination)


def convert_word_document(
    source: Path,
    output_root: Path,
    overwrite: bool = False,
    write_report: bool = False,
) -> WordConversionReport:
    source = source.expanduser().resolve()
    if not source.is_file() or source.suffix.casefold() not in WORD_SUFFIXES:
        raise WordConversionError(f"不是支持的 Word 文档：{source}")

    pandoc = find_pandoc()
    folder_name = safe_name(source.stem)
    destination = output_root.expanduser().resolve() / folder_name
    destination.parent.mkdir(parents=True, exist_ok=True)
    temp_root = destination.parent / f".{folder_name}-word-{uuid.uuid4().hex}"
    temp_root.mkdir(parents=False, exist_ok=False)
    final_temp = temp_root / folder_name
    final_temp.mkdir(parents=True)
    (final_temp / "images").mkdir()

    try:
        normalized_docx = temp_root / "normalized.docx"
        headings, table_count, _, _ = preprocess_docx(source, normalized_docx)
        lua_filter = Path(__file__).with_name("pandoc_word_filter.lua").resolve()
        if not lua_filter.is_file():
            raise WordConversionError(f"缺少 Pandoc 过滤器：{lua_filter}")
        markdown_path = final_temp / f"{folder_name}.md"
        run_checked(
            [
                str(pandoc),
                str(normalized_docx),
                "--from=docx",
                "--to=gfm",
                "--wrap=none",
                "--track-changes=accept",
                f"--lua-filter={lua_filter}",
                "--extract-media=images",
                f"--output={markdown_path.name}",
            ],
            cwd=final_temp,
        )

        markdown = markdown_path.read_text(encoding="utf-8")
        markdown, image_count = flatten_images(markdown, final_temp)
        markdown = re.sub(r"(?m)^```\s+text\s*$", "```text", markdown)
        markdown = markdown.replace("\r\n", "\n").replace("\r", "\n").rstrip() + "\n"
        converted_headings = count_markdown_headings(markdown)
        expected = {str(level): headings.get(str(level), 0) for level in range(1, 7)}
        actual = {str(level): converted_headings.get(str(level), 0) for level in range(1, 7)}
        if actual != expected:
            raise WordConversionError(
                "Word 标题样式与 Markdown 标题级别不一致："
                f"Word={expected}，Markdown={actual}"
            )
        validate_local_images(markdown, final_temp)
        markdown_path.write_text(markdown, encoding="utf-8", newline="\n")

        report = WordConversionReport(
            source=str(source),
            output=str(destination / f"{folder_name}.md"),
            pages=word_page_count(normalized_docx),
            text_characters=len(re.sub(r"\s+", "", markdown)),
            headings=headings,
            images=image_count,
            tables=table_count,
            code_blocks=count_code_blocks(markdown),
            conversion_engine=f"pandoc-{source.suffix.casefold().lstrip('.')}",
        )
        atomic_replace_dir(final_temp, destination, overwrite=overwrite)
        if write_report:
            report_dir = destination.parent / "_reports"
            report_dir.mkdir(parents=True, exist_ok=True)
            (report_dir / f"{folder_name}.json").write_text(
                json.dumps(report.as_dict(), ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
        shutil.rmtree(temp_root, ignore_errors=True)
        return report
    except Exception:
        shutil.rmtree(temp_root, ignore_errors=True)
        raise
