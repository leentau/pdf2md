from __future__ import annotations

import base64
import shutil
import unittest
import uuid
import zipfile
from pathlib import Path
from unittest.mock import patch

try:
    import pymupdf as fitz
except ImportError:
    import fitz  # type: ignore[no-redef]

from native_pdf_to_md import (
    ConversionError,
    DocumentReport,
    Element,
    Bookmark,
    code_markdown,
    convert_pdf,
    dot_leader_block_markdown,
    detect_heading,
    detect_repeated_margin_texts,
    default_output_root,
    academic_numbered_list_markdown,
    inline_markdown,
    iter_documents,
    iter_pdfs,
    join_markdown_parts,
    list_markdown,
    listing_block_markdown,
    main,
    merge_adjacent_code_parts,
    output_root_for_pdf,
    paragraph_markdown_parts,
    reference_markdown_parts,
    combine_soft_mask,
    pixmap_png_bytes,
    sort_elements_reading_order,
)
from word_to_md import convert_word_document
from mineru_pdf_to_md import (
    is_rfc2544_fake_ip,
    merge_result_archives,
    public_ipv4_answers,
    split_pdf,
)


class NativePdfToMarkdownTests(unittest.TestCase):
    def setUp(self) -> None:
        workspace = Path(__file__).resolve().parent.parent
        self.root = workspace / f"_native_pdf_test_{uuid.uuid4().hex}"
        self.root.mkdir()

    def tearDown(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)

    def make_structured_pdf(self) -> Path:
        path = self.root / "manual.pdf"
        doc = fitz.open()
        cover = doc.new_page()
        cover.insert_text((60, 100), "Native Manual", fontsize=26, fontname="hebo")
        cover.insert_text((60, 145), "Version 1.0", fontsize=12, fontname="helv")

        page = doc.new_page()
        page.insert_text((60, 75), "1. Chapter", fontsize=24, fontname="hebo")
        page.insert_text((60, 120), "Section", fontsize=16, fontname="hebo")
        page.insert_text((60, 155), "A native text paragraph.", fontsize=11, fontname="helv")
        page.insert_text((60, 185), "• Native bullet", fontsize=11, fontname="helv")
        page.insert_text((75, 220), "set_context dft -rtl", fontsize=10, fontname="cour")

        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 30), False)
        pix.clear_with(0x23A9E4)
        page.insert_text((60, 265), "Figure 1-1: Native image", fontsize=11, fontname="hebo")
        page.insert_image(fitz.Rect(60, 280, 180, 370), stream=pix.tobytes("png"))

        # A vector-only figure is rendered to images/ using its nearby caption.
        page.draw_rect(fitz.Rect(310, 285, 500, 355), color=(0, 0, 0), width=0.8)
        for x, height in ((330, 25), (360, 45), (390, 35), (420, 55), (450, 30)):
            page.draw_rect(fitz.Rect(x, 350 - height, x + 12, 350), color=(0.1, 0.4, 0.8), fill=(0.1, 0.4, 0.8))
        page.insert_text((310, 378), "Fig. 1. Vector chart", fontsize=10, fontname="helv")

        # A ruled table verifies that native word tokens, including command
        # underscores, survive cell extraction.
        for y in (420, 445, 475):
            page.draw_line((60, y), (520, y), color=(0, 0, 0), width=0.8)
        for x in (60, 220, 520):
            page.draw_line((x, 420), (x, 475), color=(0, 0, 0), width=0.8)
        page.insert_text((66, 438), "Command", fontsize=10, fontname="hebo")
        page.insert_text((226, 438), "Description", fontsize=10, fontname="hebo")
        page.insert_text((66, 463), "add_control_points", fontsize=10, fontname="cour")
        page.insert_text((226, 463), "Specifies user-defined points.", fontsize=10, fontname="helv")

        doc.set_toc([[1, "1. Chapter", 2], [2, "Section", 2]])
        doc.save(path)
        doc.close()
        return path

    def test_png_encoding_normalizes_color_and_resizes_soft_mask(self) -> None:
        color = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 20), False)
        color.clear_with(0x2A7FD1)
        mask = fitz.Pixmap(fitz.csGRAY, fitz.IRect(0, 0, 80, 40), False)
        mask.clear_with(0xB0)

        combined = combine_soft_mask(color, mask)
        payload = pixmap_png_bytes(combined)

        self.assertEqual((combined.width, combined.height), (40, 20))
        self.assertTrue(combined.alpha)
        self.assertEqual(payload[:8], b"\x89PNG\r\n\x1a\n")

    def test_repeated_running_headers_footers_and_page_numbers_are_removed(self) -> None:
        source = self.root / "running-margins.pdf"
        doc = fitz.open()
        for number in range(1, 7):
            page = doc.new_page(width=612, height=792)
            if number > 1:
                page.insert_text((485, 52), "Feedback", fontsize=8)
                page.insert_text((72, 70), "Chapter 2: Running Header", fontsize=9)
                page.insert_text((72, 716), "Example Product User Guide W-2024.09", fontsize=8)
                page.insert_text((530, 716), str(number), fontsize=8)
            page.insert_text((72, 140), f"Unique body paragraph {number}.", fontsize=11)
        doc.save(source)
        doc.close()

        with fitz.open(source) as check_doc:
            repeated = detect_repeated_margin_texts(check_doc)
        self.assertIn("feedback", repeated)
        self.assertIn("chapter 2: running header", repeated)

        convert_pdf(source, self.root / "margin-output", overwrite=True)
        markdown = (
            self.root / "margin-output" / "running-margins" / "running-margins.md"
        ).read_text(encoding="utf-8")
        self.assertNotIn("Feedback", markdown)
        self.assertNotIn("Chapter 2: Running Header", markdown)
        self.assertNotIn("Example Product User Guide", markdown)
        for number in range(1, 7):
            self.assertIn(f"Unique body paragraph {number}.", markdown)

    def test_single_and_directory_inputs_follow_path_mapping_contract(self) -> None:
        source_root = self.root / "source"
        nested = source_root / "section" / "deep"
        nested.mkdir(parents=True)
        top_pdf = source_root / "top.pdf"
        nested_pdf = nested / "nested.PDF"
        top_pdf.write_bytes(b"%PDF-test")
        nested_pdf.write_bytes(b"%PDF-test")
        word_docx = source_root / "word.docx"
        word_docx.write_bytes(b"word-test")

        discovered = [path.relative_to(source_root).as_posix() for path in iter_pdfs(source_root)]
        self.assertEqual(discovered, ["section/deep/nested.PDF", "top.pdf"])
        all_discovered = [
            path.relative_to(source_root).as_posix() for path in iter_documents(source_root)
        ]
        self.assertEqual(
            all_discovered,
            ["section/deep/nested.PDF", "top.pdf", "word.docx"],
        )

        batch_root = self.root / "converted"
        self.assertEqual(output_root_for_pdf(source_root, top_pdf, batch_root), batch_root)
        self.assertEqual(
            output_root_for_pdf(source_root, nested_pdf, batch_root),
            batch_root / "section" / "deep",
        )
        self.assertEqual(
            output_root_for_pdf(top_pdf, top_pdf, self.root / "ignored"),
            self.root / "ignored",
        )
        self.assertEqual(default_output_root(top_pdf), source_root)
        self.assertEqual(default_output_root(source_root), self.root / "source_md")
        self.assertEqual(
            output_root_for_pdf(top_pdf, top_pdf, default_output_root(top_pdf)),
            source_root,
        )
        self.assertEqual(
            output_root_for_pdf(source_root, nested_pdf, default_output_root(source_root)),
            self.root / "source_md" / "section" / "deep",
        )

    def test_directory_rerun_skips_completed_documents_without_overwrite(self) -> None:
        source_root = self.root / "manuals"
        source_root.mkdir()
        source = source_root / "done.pdf"
        source.write_bytes(b"placeholder")
        completed = self.root / "manuals_md" / "done"
        completed.mkdir(parents=True)
        (completed / "done.md").write_text("# Completed\n", encoding="utf-8")

        with patch("native_pdf_to_md.convert_pdf") as converter:
            exit_code = main([str(source_root), "--route", "native"])

        self.assertEqual(exit_code, 0)
        converter.assert_not_called()

    def test_docx_headings_tables_images_and_adjacent_code_are_preserved(self) -> None:
        from docx import Document
        from docx.shared import Inches

        source = self.root / "word_manual.docx"
        image = self.root / "pixel.png"
        image.write_bytes(
            base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
            )
        )
        document = Document()
        document.add_heading("Word Manual", level=1)
        document.add_heading("Exact Subsection", level=2)
        document.add_paragraph("Normal body text.")
        document.add_paragraph("This sentence continues to detect", style="Body Text")
        document.add_paragraph("", style="Body Text")
        document.add_paragraph('"crossings" correctly.', style="Body Text")
        document.add_paragraph("An element of one local", style="Body Text")
        document.add_paragraph("IJTAG network continues here.", style="Body Text")
        command = document.add_paragraph(style="Body Text")
        command.add_run(
            'Example "configure_scan target -type type '
            '-var_bits var_bits [\u2212var_length '
        )
        length = command.add_run("length")
        length.italic = True
        command.add_run("] -pin pin [\u2011relative_tester_cycles ")
        cycles = command.add_run("cycles")
        cycles.italic = True
        command.add_run('] [-inversion inversion]".')
        for text in ("if enabled {", "    run_test;", "}"):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(text)
            run.font.name = "Courier New"
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Command"
        table.cell(0, 1).text = "Description"
        table.cell(1, 0).text = "run_test"
        table.cell(1, 1).text = "Runs the test"
        document.add_picture(str(image), width=Inches(0.2))
        document.save(source)

        report = convert_word_document(source, self.root / "word_output")
        folder = self.root / "word_output" / "word_manual"
        markdown = (folder / "word_manual.md").read_text(encoding="utf-8")
        self.assertEqual(report.headings, {"1": 1, "2": 1})
        self.assertEqual(report.tables, 1)
        self.assertEqual(report.images, 1)
        self.assertIn("# Word Manual", markdown)
        self.assertIn("## Exact Subsection", markdown)
        self.assertIn('This sentence continues to detect "crossings" correctly.', markdown)
        self.assertIn("An element of one local IJTAG network continues here.", markdown)
        self.assertNotIn("to detect\n\n", markdown)
        self.assertNotIn("one local\n\nIJTAG", markdown)
        self.assertIn("-var_length length", markdown)
        self.assertIn("-relative_tester_cycles cycles", markdown)
        self.assertIn("configure_scan target", markdown)
        self.assertNotIn("*length*", markdown)
        self.assertNotIn("*cycles*", markdown)
        self.assertNotIn("\u2212var_length", markdown)
        self.assertNotIn("\u2011relative_tester_cycles", markdown)
        self.assertIn("```text\nif enabled {\n    run_test;\n}\n```", markdown)
        self.assertRegex(markdown, r'(?:!\[[^\]]*\]\(images/|<img[^>]+src="images/)')
        self.assertEqual(len(list((folder / "images").iterdir())), 1)

    def test_heading_levels_images_and_folder_contract(self) -> None:
        source = self.make_structured_pdf()
        output_root = self.root / "output"
        report = convert_pdf(source, output_root, write_report=True)
        folder = output_root / "manual"
        markdown = (folder / "manual.md").read_text(encoding="utf-8")

        self.assertIn("# Native Manual", markdown)
        self.assertIn("# 1. Chapter", markdown)
        self.assertIn("## Section", markdown)
        self.assertIn("- Native bullet", markdown)
        self.assertIn("```tcl", markdown)
        self.assertIn("![Figure 1-1](images/page-002-image-01.png)", markdown)
        self.assertIn("| add_control_points | Specifies user-defined points. |", markdown)
        self.assertTrue((folder / "images" / "page-002-image-01.png").is_file())
        self.assertTrue((folder / "images" / "page-002-figure-01.png").is_file())
        self.assertEqual(sorted(item.name for item in folder.iterdir()), ["images", "manual.md"])
        self.assertTrue((output_root / "_reports" / "manual.json").is_file())
        self.assertEqual(report.bookmarks_matched, 2)
        self.assertEqual(report.bookmarks_total, 2)

    def test_image_only_pdf_is_rejected_when_ocr_is_disabled(self) -> None:
        source = self.root / "scan.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.draw_rect(fitz.Rect(20, 20, 200, 200), fill=(0.2, 0.3, 0.4))
        doc.save(source)
        doc.close()

        with self.assertRaisesRegex(ConversionError, "OCR"):
            convert_pdf(source, self.root / "output")

    def test_full_page_scan_with_hidden_ocr_is_not_treated_as_native_text(self) -> None:
        source = self.root / "hidden-ocr-scan.pdf"
        doc = fitz.open()
        page = doc.new_page(width=300, height=400)
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 300, 400), False)
        pix.clear_with(0xE8E8E8)
        page.insert_image(page.rect, stream=pix.tobytes("png"))
        page.insert_text(
            (25, 60),
            "Hidden OCR text " * 12,
            fontsize=9,
            fontname="helv",
            render_mode=3,
        )
        doc.save(source)
        doc.close()

        with self.assertRaisesRegex(ConversionError, "隐藏 OCR"):
            convert_pdf(source, self.root / "output")

    def test_hidden_ocr_scan_is_automatically_routed_to_mineru(self) -> None:
        source = self.root / "automatic-ocr.pdf"
        doc = fitz.open()
        page = doc.new_page(width=300, height=400)
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 300, 400), False)
        pix.clear_with(0xF0F0F0)
        page.insert_image(page.rect, stream=pix.tobytes("png"))
        page.insert_text(
            (25, 60),
            "Existing hidden OCR text " * 12,
            fontsize=9,
            fontname="helv",
            render_mode=3,
        )
        doc.save(source)
        doc.close()

        expected = DocumentReport(
            source=str(source),
            output=str(self.root / "output" / "automatic-ocr" / "automatic-ocr.md"),
            pages=1,
            ocr_used=True,
            conversion_engine="mineru-vlm",
        )
        with patch("native_pdf_to_md.convert_pdf_with_mineru", return_value=expected) as fallback:
            report = convert_pdf(
                source,
                self.root / "output",
                mineru_token="test-token",
            )

        self.assertIs(report, expected)
        fallback.assert_called_once()
        self.assertIn("隐藏 OCR", fallback.call_args.kwargs["native_text_rejection"])
        self.assertTrue(report.as_dict()["ocr_used"])

    def test_native_pdf_can_be_forced_through_mineru_ocr(self) -> None:
        source = self.make_structured_pdf()
        expected = DocumentReport(
            source=str(source),
            output="forced.md",
            pages=2,
            ocr_used=True,
            conversion_engine="mineru-vlm",
        )
        with patch("native_pdf_to_md.convert_pdf_with_mineru", return_value=expected) as fallback:
            report = convert_pdf(
                source,
                self.root / "output",
                mineru_token="test-token",
                route="mineru-ocr",
            )
        self.assertIs(report, expected)
        self.assertEqual(
            fallback.call_args.kwargs["native_text_rejection"],
            "用户强制使用 mineru-ocr 路由",
        )

    def test_large_ocr_route_splits_and_merges_without_artificial_rule(self) -> None:
        source = self.root / "large.pdf"
        doc = fitz.open()
        for page_number in range(7):
            page = doc.new_page()
            page.insert_text((60, 100), f"Page {page_number + 1}")
        doc.save(source)
        doc.close()

        chunks = split_pdf(
            source,
            self.root / "chunks",
            max_bytes=10 * 1024 * 1024,
            max_pages=3,
        )
        chunk_page_counts = []
        for chunk in chunks:
            with fitz.open(chunk) as chunk_doc:
                chunk_page_counts.append(chunk_doc.page_count)
        self.assertEqual(chunk_page_counts, [3, 3, 1])

        archives = []
        for index in range(1, 4):
            archive = self.root / f"result-{index}.zip"
            with zipfile.ZipFile(archive, "w") as result:
                result.writestr(f"part-{index}/full.md", f"## Part {index}\n")
            archives.append(archive)
        output_dir = self.root / "merged"
        merged_path = merge_result_archives(
            archives,
            output_dir,
            output_markdown_name="large.md",
        )
        merged = merged_path.read_text(encoding="utf-8")
        self.assertEqual(merged.count("## Part"), 3)
        self.assertNotIn("\n---\n", merged)

    def test_mineru_doh_fallback_rejects_fake_and_private_addresses(self) -> None:
        payload = {
            "Answer": [
                {"type": 1, "data": "198.18.1.9"},
                {"type": 1, "data": "192.168.1.2"},
                {"type": 5, "data": "cdn.example.test"},
                {"type": 1, "data": "60.188.87.140"},
            ]
        }
        self.assertEqual(public_ipv4_answers(payload), ["60.188.87.140"])
        self.assertTrue(is_rfc2544_fake_ip("198.18.1.9"))
        self.assertFalse(is_rfc2544_fake_ip("60.188.87.140"))

    def test_toc_lines_and_code_indentation_are_preserved(self) -> None:
        toc_block = {
            "bbox": (61, 100, 550, 160),
            "lines": [
                {"bbox": (61, 100, 500, 112), "spans": [{"text": "1. Chapter........1-1"}]},
                {"bbox": (91, 120, 500, 132), "spans": [{"text": "Section........1-2"}]},
                {"bbox": (121, 140, 500, 152), "spans": [{"text": "Topic........1-3"}]},
            ],
        }
        self.assertEqual(
            listing_block_markdown(toc_block, fitz.Rect(0, 0, 612, 792)),
            "- 1. Chapter........1-1\n  - Section........1-2\n    - Topic........1-3",
        )

        code_block = {
            "bbox": (90, 100, 400, 150),
            "lines": [
                {"bbox": (90, 100, 300, 112), "spans": [{"text": "Root {"}]},
                {"bbox": (90, 120, 300, 132), "spans": [{"text": "    child : value;"}]},
                {"bbox": (90, 140, 300, 152), "spans": [{"text": "}"}]},
            ],
        }
        self.assertIn("\n    child : value;\n", code_markdown(code_block))
        self.assertEqual(join_markdown_parts(["first", "  - nested"]), "first\n\n  - nested\n")
        self.assertEqual(
            join_markdown_parts([r"complexity of $\sqrt$", r"$N$ [34]."]),
            "complexity of $\\sqrt{N}$ [34].\n",
        )

        mini_toc = {
            "bbox": (61, 100, 550, 140),
            "lines": [
                {"bbox": (61, 100, 500, 112), "spans": [{"text": "First Topic........1-1"}]},
                {"bbox": (61, 120, 500, 132), "spans": [{"text": "Second Topic........1-2"}]},
            ],
        }
        self.assertEqual(
            dot_leader_block_markdown(mini_toc),
            "- **First Topic........1-1**\n- **Second Topic........1-2**",
        )

        hierarchical = [
            list_markdown("1. Parent", level=0),
            list_markdown("a. Child", level=1),
            list_markdown("• Grandchild", level=2),
        ]
        self.assertEqual(
            "\n".join(item for item in hierarchical if item),
            "1. Parent\n    - a. Child\n        - Grandchild",
        )

        chapter_block = {
            "bbox": (60, 100, 300, 130),
            "lines": [{"spans": [{"text": "Introduction to PrimeTime", "font": "Bold", "size": 22}]}],
        }
        chapter = Bookmark(
            level=1,
            title="1 Introduction to PrimeTime",
            page=38,
            key="1 introduction to primetime",
        )
        level, matched_bookmark = detect_heading(
            "Introduction to PrimeTime",
            chapter_block,
            [chapter],
            set(),
            38,
        )
        self.assertEqual(level, 1)
        self.assertIs(matched_bookmark, chapter)

        self.assertEqual(
            merge_adjacent_code_parts(
                [
                    "```text\nfirst\n```",
                    "<!-- PDF page 2 -->",
                    "```tcl\n  second\n```",
                    "ordinary paragraph",
                ]
            ),
            ["```text\nfirst\n  second\n```", "ordinary paragraph"],
        )

        page_rect = fitz.Rect(0, 0, 612, 792)
        title = Element("text", fitz.Rect(60, 40, 550, 80), "title")
        right_first = Element("text", fitz.Rect(312, 100, 563, 160), "right-first")
        left_first = Element("text", fitz.Rect(49, 110, 300, 180), "left-first")
        left_second = Element("text", fitz.Rect(49, 200, 300, 260), "left-second")
        right_second = Element("text", fitz.Rect(312, 210, 563, 270), "right-second")
        ordered = sort_elements_reading_order(
            [right_first, left_first, title, right_second, left_second], page_rect
        )
        self.assertEqual(
            [element.data for element in ordered],
            ["title", "left-first", "left-second", "right-first", "right-second"],
        )

    def test_academic_math_indent_references_and_numbered_items(self) -> None:
        math_block = {
            "bbox": (49, 100, 200, 112),
            "lines": [{
                "bbox": (49, 100, 200, 112),
                "spans": [
                    {"text": "variance ", "font": "Times-Roman", "size": 10, "origin": (49, 110)},
                    {"text": "λ", "font": "CMMI10", "size": 10, "origin": (100, 110)},
                    {"text": "j", "font": "CMMI7", "size": 7, "origin": (106, 111)},
                ],
            }],
        }
        self.assertEqual(inline_markdown(math_block), "variance $\\lambda_{j}$")

        set_block = {
            "bbox": (49, 100, 150, 112),
            "lines": [{
                "bbox": (49, 100, 150, 112),
                "spans": [
                    {"text": "i", "font": "CMMI7", "size": 7, "origin": (49, 111)},
                    {"text": "∈", "font": "CMSY7", "size": 7, "origin": (55, 111)},
                    {"text": "S", "font": "CMMI7", "size": 7, "origin": (61, 111)},
                    {"text": " π", "font": "CMMI10", "size": 10, "origin": (70, 110)},
                    {"text": "i", "font": "CMMI7", "size": 7, "origin": (77, 111)},
                ],
            }],
        }
        self.assertEqual(inline_markdown(set_block), "$_{i\\in S}\\pi_{i}$")

        styled_block = {
            "bbox": (49, 100, 200, 132),
            "lines": [
                {"bbox": (49, 100, 200, 112), "spans": [{"text": "first", "font": "Bold", "flags": 16}]},
                {"bbox": (49, 120, 200, 132), "spans": [{"text": "second", "font": "Bold", "flags": 16}]},
            ],
        }
        self.assertEqual(inline_markdown(styled_block), "**first second**")

        baseline_bold_block = {
            "bbox": (49, 100, 300, 132),
            "lines": [{
                "bbox": (49, 100, 300, 112),
                "spans": [{"text": "A" * 90, "font": "Medium", "flags": 16}],
            }],
        }
        self.assertEqual(inline_markdown(baseline_bold_block), "A" * 90)

        contact_block = {
            "bbox": (49, 100, 300, 132),
            "lines": [
                {"bbox": (49, 100, 200, 112), "spans": [{"text": "person@example.com", "font": "Italic", "flags": 2}]},
                {"bbox": (49, 120, 300, 132), "spans": [
                    {"text": "{", "font": "CMSY10", "flags": 6},
                    {"text": "one, two", "font": "Italic", "flags": 2},
                    {"text": "}@example.com", "font": "Italic", "flags": 2},
                ]},
            ],
        }
        contact = inline_markdown(contact_block)
        self.assertIn("person@example.com", contact)
        self.assertIn("one@example.com, two@example.com", contact)
        self.assertNotIn("{", contact)
        self.assertNotIn("}", contact)
        self.assertNotIn("$", contact)
        self.assertNotIn("*", contact)

        paragraph_block = {
            "bbox": (49, 100, 300, 140),
            "lines": [
                {"bbox": (49, 100, 300, 112), "spans": [{"text": "first paragraph"}]},
                {"bbox": (59, 120, 300, 132), "spans": [{"text": "second paragraph"}]},
            ],
        }
        self.assertEqual(
            paragraph_markdown_parts(paragraph_block, fitz.Rect(0, 0, 612, 792), True),
            ["first paragraph", "&emsp;&emsp;second paragraph"],
        )

        reference_block = {
            "bbox": (49, 100, 300, 160),
            "lines": [
                {"bbox": (49, 100, 300, 112), "spans": [{"text": "[1] First reference."}]},
                {"bbox": (49, 120, 300, 132), "spans": [{"text": "[2] Second"}]},
                {"bbox": (49, 140, 300, 152), "spans": [{"text": "reference."}]},
            ],
        }
        self.assertEqual(reference_markdown_parts(reference_block), ["[1] First reference.", "[2] Second reference."])

        numbered_block = {
            "bbox": (49, 100, 300, 160),
            "lines": [
                {"bbox": (49, 100, 300, 112), "spans": [{"text": "1) First item"}]},
                {"bbox": (64, 120, 300, 132), "spans": [{"text": "continues"}]},
                {"bbox": (49, 140, 300, 152), "spans": [{"text": "2) Second item"}]},
            ],
        }
        self.assertEqual(academic_numbered_list_markdown(numbered_block), "1. First item continues\n2. Second item")


if __name__ == "__main__":
    unittest.main()
